"""Populate the supplied Word confirmation letter template from a verified answer.

The template's merge fields are plain `{{ FIELD_NAME }}` placeholders, each
sitting in its own docx run (confirmed by inspection -- python-docx did not
split any of them across runs), so no run-merging workaround is needed: find
the run, replace its text, clear the yellow highlight that marked it as
unfilled.
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import docx

_PLACEHOLDER_RE = re.compile(r"^\{\{\s*([A-Z_]+)\s*\}\}$")
_ANY_PLACEHOLDER_RE = re.compile(r"\{\{[^}]*\}\}")

_SOURCE_FIELDS = (
    "destination_region_sections",
    "fare_era_sections",
    "checked_baggage_sections",
    "carry_on_sections",
)


def _collect_source_section_numbers(answer: dict) -> list[str]:
    seen: list[str] = []
    for field in _SOURCE_FIELDS:
        for num in re.findall(r"\d+", answer.get(field, "")):
            if num not in seen:
                seen.append(num)
    return seen


def _source_reference(answer: dict, processed_path: Path) -> str:
    parsed = json.loads(processed_path.read_text())
    titles = {s["section_number"]: s["title"] for s in parsed["sections"] if s["section_number"]}
    numbers = _collect_source_section_numbers(answer)
    return "; ".join(f"Section {n}: {titles[n]}" for n in numbers if n in titles)


def build_merge_fields(answer: dict, processed_path: Path, *, today: date | None = None) -> dict[str, str]:
    today = today or date.today()
    return {
        "DATE": f"{today.day} {today:%B %Y}",
        "TICKET_ISSUE_DATE": answer["ticket_issue_date"],
        "FARE_TYPE": answer["fare_type"],
        "DESTINATION_REGION": answer["destination_region"],
        "CHECKED_BAGGAGE_ALLOWANCE": answer["checked_baggage_allowance"],
        "CARRY_ON_ALLOWANCE": answer["carry_on_allowance"],
        "SOURCE_REFERENCE": _source_reference(answer, processed_path),
    }


def _fill_paragraph(paragraph, fields: dict[str, str]) -> None:
    for run in paragraph.runs:
        match = _PLACEHOLDER_RE.match(run.text.strip())
        if match and match.group(1) in fields:
            run.text = fields[match.group(1)]
            run.font.highlight_color = None


def _find_unresolved_placeholders(document) -> list[str]:
    found: list[str] = []
    for paragraph in document.paragraphs:
        found.extend(_ANY_PLACEHOLDER_RE.findall(paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                found.extend(_ANY_PLACEHOLDER_RE.findall(cell.text))
    return found


def populate_letter(template_path: Path, fields: dict[str, str], output_path: Path) -> Path:
    document = docx.Document(str(template_path))

    for paragraph in document.paragraphs:
        _fill_paragraph(paragraph, fields)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _fill_paragraph(paragraph, fields)

    unresolved = _find_unresolved_placeholders(document)
    if unresolved:
        raise ValueError(f"Unresolved placeholders remain in output: {unresolved}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
