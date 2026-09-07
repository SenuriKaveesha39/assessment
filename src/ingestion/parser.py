"""Parse the Conditions of Carriage .docx into semantic sections with metadata.

Design notes
------------
- We walk the document body in original XML order (`_iter_block_items`)
  rather than using `Document.paragraphs` / `Document.tables` separately,
  because those two collections lose each other's relative position -- and
  in this document a table (international baggage allowances by region) is
  nested *inside* a section, between two paragraphs. Using the flattened
  collections would silently detach the table from the section that gives
  it meaning.
- Sections are delimited by heading paragraphs (any "Heading *" style below
  the document's own Heading 1 title). Everything until the next heading
  belongs to that section.
- Each section is tagged with `service_scope` (domestic/international) and
  `fare_era` (current/legacy) inferred from its own title + body text. This
  directly targets the failure mode called out in the brief: a RAG system
  confusing domestic vs. international allowances, or the wrong fare era.
  The tags let a retriever filter by metadata *in addition to* semantic
  similarity, rather than relying on embeddings alone to keep these apart.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table as _DocxTable
from docx.text.paragraph import Paragraph

from .models import DocumentMetadata, ParsedDocument, Section, Table, TableRow

_SECTION_NUMBER_RE = re.compile(r"^(\d+)\.\s*(.+)$")
_VERSION_RE = re.compile(
    r"Document Version:\s*(?P<version>[\d.]+)\s*\(Effective\s*(?P<date>[^)]+)\)",
    re.IGNORECASE,
)


def _iter_block_items(document: _Document):
    """Yield paragraphs and tables from the document body, in document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield _DocxTable(child, document)


def _heading_level(style_name: str | None) -> int | None:
    if not style_name or not style_name.startswith("Heading"):
        return None
    try:
        return int(style_name.rsplit(" ", 1)[-1])
    except ValueError:
        return None


def _table_to_model(table: _DocxTable) -> Table:
    rows = list(table.rows)
    headers = [c.text.strip() for c in rows[0].cells] if rows else []
    data_rows = []
    for row in rows[1:]:
        values = [c.text.strip() for c in row.cells]
        data_rows.append(TableRow(cells=dict(zip(headers, values))))
    return Table(headers=headers, rows=data_rows)


def _infer_service_scope(title: str, body_text: str) -> list[str]:
    # A title that names exactly one scope (e.g. "... (Domestic, Current
    # Fares)") is authoritative for that section, even if its body text goes
    # on to cross-reference the other scope (as Section 5 does when pointing
    # readers to Section 6 for international allowances). Falling back to a
    # body-text scan for such mentions would wrongly tag the section as
    # applying to both.
    title_lower = title.lower()
    has_domestic_title = "domestic" in title_lower
    has_international_title = "international" in title_lower
    if has_domestic_title != has_international_title:
        return ["domestic"] if has_domestic_title else ["international"]

    text = body_text.lower()
    scope: list[str] = []
    if "international" in text:
        scope.append("international")
    if "domestic" in text:
        scope.append("domestic")
    return scope


def _infer_fare_era(combined_text: str) -> str:
    text = combined_text.lower()
    if "legacy" in text or "issued before" in text:
        return "legacy"
    if "current fares" in text or "issued on or after" in text:
        return "current"
    return "unspecified"


def _extract_document_metadata(
    document: _Document, *, source_url: str, retrieved_at: str, sha256: str, size_bytes: int
) -> DocumentMetadata:
    title = "Conditions of Carriage"
    version = None
    effective_date = None

    for block in _iter_block_items(document):
        if not isinstance(block, Paragraph):
            continue
        style = block.style.name if block.style else None
        text = block.text.strip()
        if not text:
            continue
        if _heading_level(style) == 1:
            title = text
        match = _VERSION_RE.search(text)
        if match:
            version = match.group("version")
            effective_date = match.group("date").strip()
        if version is not None:
            break  # title + version always precede the first numbered section

    return DocumentMetadata(
        title=title,
        version=version,
        effective_date=effective_date,
        source_url=source_url,
        retrieved_at=retrieved_at,
        sha256=sha256,
        size_bytes=size_bytes,
    )


def parse_document(
    path: Path,
    *,
    source_url: str,
    retrieved_at: str,
    sha256: str,
    size_bytes: int,
) -> ParsedDocument:
    document = docx.Document(str(path))
    metadata = _extract_document_metadata(
        document,
        source_url=source_url,
        retrieved_at=retrieved_at,
        sha256=sha256,
        size_bytes=size_bytes,
    )

    sections: list[Section] = []
    current: Section | None = None
    doc_title_seen = False

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else None
            level = _heading_level(style)
            text = block.text.strip()
            if not text:
                continue

            if level == 1:
                doc_title_seen = True
                continue

            if doc_title_seen and level is not None:
                # Start of a new numbered section.
                match = _SECTION_NUMBER_RE.match(text)
                number, title = (match.group(1), match.group(2)) if match else (None, text)
                current = Section(section_number=number, title=title)
                sections.append(current)
                continue

            if current is None:
                continue  # front-matter before the first numbered section

            if style == "List Bullet":
                current.bullets.append(text)
            else:
                current.paragraphs.append(text)
        else:  # table
            if current is not None:
                current.tables.append(_table_to_model(block))

    for section in sections:
        body = " ".join([*section.paragraphs, *section.bullets, *(t.to_text() for t in section.tables)])
        combined = f"{section.title} {body}"
        section.service_scope = _infer_service_scope(section.title, body)
        section.fare_era = _infer_fare_era(combined)

    return ParsedDocument(metadata=metadata, sections=sections)
